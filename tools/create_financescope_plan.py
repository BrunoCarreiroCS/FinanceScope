from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCX_PATH = DOCS / "FinanceScope_Plano_Desenvolvimento.docx"
PDF_PATH = DOCS / "FinanceScope_Plano_Desenvolvimento.pdf"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_FILL = "F2F4F7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.bold = bold
    if bold:
        run.font.color.rgb = RGBColor.from_string(INK)


def add_docx_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, LIGHT_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else DARK_BLUE)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("FinanceScope")
    run.font.name = "Calibri"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Plano tecnico e mapa visual de desenvolvimento")
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Projeto beta para GitHub, entrevistas e evolucao tecnica | Junho de 2026"
    ).italic = True

    doc.add_page_break()

    add_heading(doc, "1. Visao geral", 1)
    add_para(
        doc,
        "FinanceScope e um sistema web desktop-first de gestao financeira pessoal "
        "com foco em dashboard, diagnostico e simulacao de decisoes. O objetivo do "
        "projeto beta e provar o diferencial RealCost: traduzir gastos em horas de "
        "trabalho, impacto na renda, risco de fechamento do mes e atraso em metas.",
    )
    add_para(
        doc,
        "O nome FinanceScope deve ser usado no repositorio, no README, nos prints, "
        "nos commits de apresentacao e neste plano. O nome final do site/produto "
        "permanece como decisao futura, sem bloquear o MVP.",
    )
    add_docx_table(
        doc,
        ["Item", "Definicao"],
        [
            ["Produto", "Dashboard financeiro com analise de custo real"],
            ["Publico inicial", "Estudantes, estagiarios e jovens profissionais"],
            ["Stack", "Python Flask, SQLite, HTML, CSS, JavaScript e Chart.js"],
            ["Formato", "Aplicacao web desktop-first, adaptavel para notebook/tablet"],
            ["Diferencial", "RealCost Engine como motor de interpretacao financeira"],
        ],
    )

    add_heading(doc, "2. Melhorias sobre o escopo original", 1)
    add_bullets(
        doc,
        [
            "Padronizar o nome como FinanceScope para evitar divergencia entre PDF, GitHub e apresentacao.",
            "Enxugar o MVP para priorizar perfil financeiro, transacoes, dashboard, metas, simulador e RealCost Engine.",
            "Mover login, score financeiro, assinaturas avancadas, relatorios completos, importacao CSV e exportacao PDF para roadmap futuro.",
            "Tratar o dashboard como ferramenta de decisao, nao apenas como painel de graficos.",
            "Exibir formulas e explicacoes simples para aumentar confianca e demonstrar maturidade tecnica.",
            "Usar linguagem de estimativa educacional, evitando recomendacao financeira personalizada.",
            "Adicionar campos de auditoria e padronizacao no modelo de dados para facilitar evolucao.",
        ],
    )

    add_heading(doc, "3. Mapa visual das fases", 1)
    add_docx_table(
        doc,
        ["Fase", "Entrega", "Resultado esperado"],
        [
            ["1", "Base do projeto e identidade visual", "Estrutura Flask, templates, CSS base, layout desktop e navegacao principal."],
            ["2", "Perfil financeiro e dados iniciais", "Usuario define renda, horas mensais, limite do mes e meta principal."],
            ["3", "CRUD de transacoes", "Receitas e despesas podem ser criadas, editadas, filtradas e excluidas."],
            ["4", "Dashboard e graficos", "Cards, categorias, evolucao mensal, ranking de gastos e alertas basicos funcionando."],
            ["5", "RealCost Engine", "Gastos convertidos em horas, percentual da renda, risco e atraso de meta."],
            ["6", "Metas e simulador", "Usuario acompanha objetivos e testa compras antes de registrar."],
            ["7", "Polimento e GitHub", "README, prints, dados demo, testes, deploy e apresentacao final."],
        ],
    )
    add_para(
        doc,
        "Ordem recomendada: construir primeiro a base navegavel, depois persistencia, "
        "entao analise. O RealCost deve entrar quando ja houver dados suficientes "
        "para demonstrar valor real no dashboard e no simulador.",
    )

    add_heading(doc, "4. Regras funcionais do MVP", 1)
    add_docx_table(
        doc,
        ["Modulo", "Regras obrigatorias"],
        [
            ["Perfil financeiro", "Guardar renda mensal, horas mensais trabalhadas/estudadas, limite planejado e meta principal."],
            ["Transacoes", "Permitir tipo receita/despesa, descricao, valor positivo, data, categoria, forma de pagamento e recorrencia simples."],
            ["Dashboard", "Mostrar receitas, despesas, saldo, previsao de fechamento, maior vazamento, graficos e alertas."],
            ["Metas", "Criar meta com valor alvo, valor atual e aporte mensal; calcular progresso e meses estimados."],
            ["Simulador", "Receber item, valor, categoria, prioridade e parcelas; retornar impacto antes da decisao."],
            ["RealCost", "Centralizar formulas em utilitario separado e reutilizar no dashboard, simulador e detalhes de despesa."],
            ["Alertas", "Classificar risco como baixo, medio ou alto com explicacao curta e nao proibitiva."],
        ],
    )

    add_heading(doc, "5. Regras de calculo", 1)
    add_docx_table(
        doc,
        ["Calculo", "Formula", "Uso no produto"],
        [
            ["Valor da hora", "renda_mensal / horas_mensais", "Base para traduzir compras em tempo de esforco."],
            ["Custo em horas", "valor / valor_da_hora", "Exibir em horas e minutos no simulador e nos detalhes."],
            ["Impacto na renda", "(valor / renda_mensal) * 100", "Mostrar peso percentual de compras e categorias."],
            ["Saldo do mes", "total_receitas - total_despesas", "Card principal do dashboard."],
            ["Previsao de fechamento", "media_diaria_despesas * dias_do_mes", "Alerta de risco quando ultrapassar limite planejado."],
            ["Tempo da meta", "(valor_alvo - valor_atual) / aporte_mensal", "Estimativa de meses para concluir objetivo."],
            ["Atraso da meta", "valor_compra / aporte_mensal", "Quantos meses de aporte a compra consome."],
            ["Custo anual recorrente", "valor_mensal * 12", "Roadmap para modulo de assinaturas."],
        ],
    )
    add_para(
        doc,
        "Se renda mensal, horas mensais ou aporte mensal forem zero ou vazios, o "
        "sistema deve evitar divisao por zero e mostrar uma mensagem orientando o "
        "usuario a completar o perfil financeiro.",
    )

    add_heading(doc, "6. Arquitetura tecnica", 1)
    add_docx_table(
        doc,
        ["Camada", "Responsabilidade"],
        [
            ["Templates HTML", "Estruturar telas, formularios, tabelas e componentes principais."],
            ["CSS", "Layout SaaS desktop-first, sidebar, cards, responsividade e estados visuais."],
            ["JavaScript", "Interacoes, validacoes leves, Chart.js e atualizacoes de interface."],
            ["Flask", "Rotas, processamento de formularios, regras de negocio e integracao com banco."],
            ["SQLite", "Persistencia local de perfil, transacoes, metas, categorias e simulacoes."],
            ["utils/finance.py", "Formulas do RealCost, score futuro e funcoes de previsao."],
            ["tests/", "Testes das formulas e dos fluxos principais."],
        ],
    )
    add_para(
        doc,
        "A primeira versao pode usar usuario unico local. Autenticacao deve entrar "
        "somente depois que o diferencial analitico estiver funcionando e apresentavel.",
    )

    add_heading(doc, "7. Modelo de dados recomendado", 1)
    add_docx_table(
        doc,
        ["Tabela", "Campos principais"],
        [
            ["users", "id, name, monthly_income, monthly_hours, monthly_limit, created_at, updated_at"],
            ["categories", "id, name, type, color, is_default, created_at"],
            ["transactions", "id, user_id, type, description, category_id, amount, date, payment_method, is_recurring, created_at, updated_at"],
            ["goals", "id, user_id, name, target_amount, current_amount, monthly_contribution, priority, created_at, updated_at"],
            ["purchase_simulations", "id, user_id, item_name, amount, category_id, installments, risk_level, result_json, created_at"],
        ],
    )
    add_para(
        doc,
        "Padronizacoes recomendadas: type deve aceitar apenas income ou expense; "
        "valores monetarios devem ser sempre positivos; recorrencia deve ser booleana "
        "no MVP e pode virar tabela propria em versao futura.",
    )

    add_heading(doc, "8. Criterios de aceite por fase", 1)
    add_docx_table(
        doc,
        ["Fase", "Aceite"],
        [
            ["1", "A aplicacao abre localmente, possui layout consistente, sidebar e telas principais navegaveis."],
            ["2", "Perfil salva renda, horas mensais, limite e meta principal; validacoes impedem valores invalidos."],
            ["3", "CRUD de transacoes persiste em SQLite e filtros por mes, categoria e tipo funcionam."],
            ["4", "Dashboard exibe cards corretos, dois graficos Chart.js e ranking dos maiores gastos."],
            ["5", "RealCost calcula horas, impacto percentual, risco e atraso de meta com testes unitarios."],
            ["6", "Metas e simulador usam dados reais do perfil e mostram mensagens claras de decisao."],
            ["7", "README completo, dados demo, prints, instrucoes de execucao, deploy e revisao final."],
        ],
    )

    add_heading(doc, "9. Checklist final para GitHub", 1)
    add_bullets(
        doc,
        [
            "README com problema, solucao, diferencial RealCost, stack, funcionalidades e prints.",
            "Secao Como rodar com ambiente virtual, dependencias, inicializacao do banco e comando Flask.",
            "Dados de demonstracao para apresentar dashboard, graficos, metas e simulador sem cadastro manual demorado.",
            "GIF curto mostrando cadastro de despesa e resultado do simulador Posso Comprar?.",
            "Aviso educacional: as analises sao estimativas baseadas nos dados informados e nao substituem planejamento financeiro profissional.",
            "Issues ou roadmap com login, importacao CSV, assinaturas completas, relatorios PDF, API externa e responsividade mobile.",
            "Commits organizados por fase para demonstrar processo de desenvolvimento.",
        ],
    )

    add_heading(doc, "10. Backlog recomendado", 1)
    add_docx_table(
        doc,
        ["Prioridade", "Item", "Motivo"],
        [
            ["MVP", "Perfil, transacoes, dashboard, RealCost, metas e simulador", "Prova o diferencial com menor escopo viavel."],
            ["Pos-MVP", "Score financeiro e assinaturas completas", "Aumenta profundidade apos dados centrais estarem estaveis."],
            ["Pos-MVP", "Relatorios por periodo e exportacao PDF", "Ajuda apresentacao e evolucao tecnica."],
            ["Futuro", "Login e multiplos usuarios", "Necessario para produto real, mas nao para demonstracao inicial."],
            ["Futuro", "Importacao CSV e integracao bancaria", "Eleva complexidade e deve vir apenas depois da base validada."],
        ],
    )

    add_heading(doc, "11. Regra de comunicacao do produto", 1)
    add_para(
        doc,
        "O FinanceScope deve ajudar o usuario a decidir, nao mandar nele. As mensagens "
        "devem ser objetivas, explicaveis e transparentes. Exemplo: 'Esta compra "
        "equivale a 18h40 de trabalho e pode atrasar sua meta em 1,2 mes, considerando "
        "os dados informados no perfil.'",
    )
    add_bullets(
        doc,
        [
            "Use 'estimativa', 'pode indicar', 'com base nos seus dados' e 'no ritmo atual'.",
            "Evite frases como 'nao compre', 'voce deve investir' ou 'garantia de economia'.",
            "Mostre sempre a formula ou uma explicacao simples perto do resultado principal.",
        ],
    )

    add_heading(doc, "12. Definicao de pronto do MVP", 1)
    add_numbered(
        doc,
        [
            "Usuario configura perfil financeiro e visualiza dados no dashboard.",
            "Usuario cadastra, edita, filtra e remove transacoes.",
            "Dashboard mostra receitas, despesas, saldo, categorias, evolucao e alertas.",
            "RealCost Engine tem testes automatizados para os calculos principais.",
            "Usuario cria meta e o sistema calcula progresso e prazo estimado.",
            "Simulador mostra impacto de compra em horas, renda, risco e meta.",
            "Projeto possui README, dados demo, prints, aviso educacional e instrucoes de execucao.",
        ],
    )

    doc.core_properties.title = "FinanceScope - Plano de Desenvolvimento"
    doc.core_properties.author = "Bruno Carreiro dos Santos"
    doc.core_properties.subject = "Plano tecnico e mapa visual do projeto FinanceScope"
    doc.save(DOCX_PATH)


def styles():
    base = getSampleStyleSheet()
    base.add(
        ParagraphStyle(
            "CoverTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=28,
            textColor=colors.HexColor(f"#{INK}"),
            alignment=TA_CENTER,
            spaceAfter=8,
        )
    )
    base.add(
        ParagraphStyle(
            "CoverSubtitle",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=14,
            textColor=colors.HexColor(f"#{DARK_BLUE}"),
            alignment=TA_CENTER,
            leading=18,
            spaceAfter=18,
        )
    )
    base.add(
        ParagraphStyle(
            "H1Custom",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceBefore=14,
            spaceAfter=8,
        )
    )
    base.add(
        ParagraphStyle(
            "BodyCustom",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=12,
            spaceAfter=6,
        )
    )
    base.add(
        ParagraphStyle(
            "BulletCustom",
            parent=base["BodyCustom"],
            leftIndent=14,
            firstLineIndent=-8,
            bulletIndent=0,
        )
    )
    return base


def pdf_table(data: list[list[str]], col_widths: list[float] | None = None) -> Table:
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_FILL}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D3DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def p(text: str, style_name: str = "BodyCustom") -> Paragraph:
    return Paragraph(text, PDF_STYLES[style_name])


def bullet(text: str) -> Paragraph:
    return Paragraph(f"- {text}", PDF_STYLES["BulletCustom"])


PDF_STYLES = styles()


def build_pdf() -> None:
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.72 * inch,
        bottomMargin=0.72 * inch,
        title="FinanceScope - Plano de Desenvolvimento",
        author="Bruno Carreiro dos Santos",
    )
    story = []
    story.append(Spacer(1, 1.5 * inch))
    story.append(p("FinanceScope", "CoverTitle"))
    story.append(p("Plano tecnico e mapa visual de desenvolvimento", "CoverSubtitle"))
    story.append(
        p(
            "Projeto beta para GitHub, entrevistas e evolucao tecnica<br/>Junho de 2026",
            "CoverSubtitle",
        )
    )
    story.append(PageBreak())

    sections = [
        (
            "1. Visao geral",
            [
                "FinanceScope e um sistema web desktop-first de gestao financeira pessoal com foco em dashboard, diagnostico e simulacao de decisoes. O objetivo do projeto beta e provar o diferencial RealCost: traduzir gastos em horas de trabalho, impacto na renda, risco de fechamento do mes e atraso em metas.",
                "O nome FinanceScope deve ser usado no repositorio, no README, nos prints, nos commits de apresentacao e neste plano. O nome final do site/produto permanece como decisao futura, sem bloquear o MVP.",
            ],
        ),
        (
            "2. Melhorias sobre o escopo original",
            [
                "- Padronizar o nome como FinanceScope para evitar divergencia entre PDF, GitHub e apresentacao.",
                "- Enxugar o MVP para priorizar perfil financeiro, transacoes, dashboard, metas, simulador e RealCost Engine.",
                "- Mover login, score financeiro, assinaturas avancadas, relatorios completos, importacao CSV e exportacao PDF para roadmap futuro.",
                "- Tratar o dashboard como ferramenta de decisao, nao apenas como painel de graficos.",
                "- Exibir formulas e explicacoes simples para aumentar confianca e demonstrar maturidade tecnica.",
                "- Usar linguagem de estimativa educacional, evitando recomendacao financeira personalizada.",
                "- Adicionar campos de auditoria e padronizacao no modelo de dados para facilitar evolucao.",
            ],
        ),
    ]
    for heading, paragraphs in sections:
        story.append(p(heading, "H1Custom"))
        for paragraph in paragraphs:
            story.append(bullet(paragraph[2:]) if paragraph.startswith("- ") else p(paragraph))

    story.append(p("3. Mapa visual das fases", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Fase", "Entrega", "Resultado esperado"],
                ["1", "Base do projeto e identidade visual", "Estrutura Flask, templates, CSS base, layout desktop e navegacao principal."],
                ["2", "Perfil financeiro e dados iniciais", "Usuario define renda, horas mensais, limite do mes e meta principal."],
                ["3", "CRUD de transacoes", "Receitas e despesas podem ser criadas, editadas, filtradas e excluidas."],
                ["4", "Dashboard e graficos", "Cards, categorias, evolucao mensal, ranking de gastos e alertas basicos funcionando."],
                ["5", "RealCost Engine", "Gastos convertidos em horas, percentual da renda, risco e atraso de meta."],
                ["6", "Metas e simulador", "Usuario acompanha objetivos e testa compras antes de registrar."],
                ["7", "Polimento e GitHub", "README, prints, dados demo, testes, deploy e apresentacao final."],
            ],
            [0.45 * inch, 2.1 * inch, 3.7 * inch],
        )
    )

    story.append(p("4. Regras funcionais do MVP", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Modulo", "Regras obrigatorias"],
                ["Perfil financeiro", "Guardar renda mensal, horas mensais trabalhadas/estudadas, limite planejado e meta principal."],
                ["Transacoes", "Permitir tipo receita/despesa, descricao, valor positivo, data, categoria, forma de pagamento e recorrencia simples."],
                ["Dashboard", "Mostrar receitas, despesas, saldo, previsao de fechamento, maior vazamento, graficos e alertas."],
                ["Metas", "Criar meta com valor alvo, valor atual e aporte mensal; calcular progresso e meses estimados."],
                ["Simulador", "Receber item, valor, categoria, prioridade e parcelas; retornar impacto antes da decisao."],
                ["RealCost", "Centralizar formulas em utilitario separado e reutilizar no dashboard, simulador e detalhes de despesa."],
                ["Alertas", "Classificar risco como baixo, medio ou alto com explicacao curta e nao proibitiva."],
            ],
            [1.45 * inch, 4.8 * inch],
        )
    )

    story.append(p("5. Regras de calculo", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Calculo", "Formula", "Uso no produto"],
                ["Valor da hora", "renda_mensal / horas_mensais", "Base para traduzir compras em tempo de esforco."],
                ["Custo em horas", "valor / valor_da_hora", "Exibir em horas e minutos no simulador e nos detalhes."],
                ["Impacto na renda", "(valor / renda_mensal) * 100", "Mostrar peso percentual de compras e categorias."],
                ["Saldo do mes", "total_receitas - total_despesas", "Card principal do dashboard."],
                ["Previsao de fechamento", "media_diaria_despesas * dias_do_mes", "Alerta de risco quando ultrapassar limite planejado."],
                ["Tempo da meta", "(valor_alvo - valor_atual) / aporte_mensal", "Estimativa de meses para concluir objetivo."],
                ["Atraso da meta", "valor_compra / aporte_mensal", "Quantos meses de aporte a compra consome."],
                ["Custo anual recorrente", "valor_mensal * 12", "Roadmap para modulo de assinaturas."],
            ],
            [1.35 * inch, 2.1 * inch, 2.8 * inch],
        )
    )
    story.append(
        p(
            "Se renda mensal, horas mensais ou aporte mensal forem zero ou vazios, o sistema deve evitar divisao por zero e orientar o usuario a completar o perfil financeiro."
        )
    )

    story.append(PageBreak())
    story.append(p("6. Arquitetura tecnica", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Camada", "Responsabilidade"],
                ["Templates HTML", "Estruturar telas, formularios, tabelas e componentes principais."],
                ["CSS", "Layout SaaS desktop-first, sidebar, cards, responsividade e estados visuais."],
                ["JavaScript", "Interacoes, validacoes leves, Chart.js e atualizacoes de interface."],
                ["Flask", "Rotas, processamento de formularios, regras de negocio e integracao com banco."],
                ["SQLite", "Persistencia local de perfil, transacoes, metas, categorias e simulacoes."],
                ["utils/finance.py", "Formulas do RealCost, score futuro e funcoes de previsao."],
                ["tests/", "Testes das formulas e dos fluxos principais."],
            ],
            [1.45 * inch, 4.8 * inch],
        )
    )

    story.append(p("7. Modelo de dados recomendado", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Tabela", "Campos principais"],
                ["users", "id, name, monthly_income, monthly_hours, monthly_limit, created_at, updated_at"],
                ["categories", "id, name, type, color, is_default, created_at"],
                ["transactions", "id, user_id, type, description, category_id, amount, date, payment_method, is_recurring, created_at, updated_at"],
                ["goals", "id, user_id, name, target_amount, current_amount, monthly_contribution, priority, created_at, updated_at"],
                ["purchase_simulations", "id, user_id, item_name, amount, category_id, installments, risk_level, result_json, created_at"],
            ],
            [1.35 * inch, 4.9 * inch],
        )
    )
    story.append(
        p(
            "Padronizacoes recomendadas: type deve aceitar apenas income ou expense; valores monetarios devem ser sempre positivos; recorrencia deve ser booleana no MVP e pode virar tabela propria em versao futura."
        )
    )

    story.append(p("8. Criterios de aceite por fase", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Fase", "Aceite"],
                ["1", "A aplicacao abre localmente, possui layout consistente, sidebar e telas principais navegaveis."],
                ["2", "Perfil salva renda, horas mensais, limite e meta principal; validacoes impedem valores invalidos."],
                ["3", "CRUD de transacoes persiste em SQLite e filtros por mes, categoria e tipo funcionam."],
                ["4", "Dashboard exibe cards corretos, dois graficos Chart.js e ranking dos maiores gastos."],
                ["5", "RealCost calcula horas, impacto percentual, risco e atraso de meta com testes unitarios."],
                ["6", "Metas e simulador usam dados reais do perfil e mostram mensagens claras de decisao."],
                ["7", "README completo, dados demo, prints, instrucoes de execucao, deploy e revisao final."],
            ],
            [0.45 * inch, 5.8 * inch],
        )
    )

    story.append(p("9. Checklist final para GitHub", "H1Custom"))
    for item in [
        "README com problema, solucao, diferencial RealCost, stack, funcionalidades e prints.",
        "Secao Como rodar com ambiente virtual, dependencias, inicializacao do banco e comando Flask.",
        "Dados de demonstracao para apresentar dashboard, graficos, metas e simulador sem cadastro manual demorado.",
        "GIF curto mostrando cadastro de despesa e resultado do simulador Posso Comprar?.",
        "Aviso educacional: as analises sao estimativas baseadas nos dados informados e nao substituem planejamento financeiro profissional.",
        "Issues ou roadmap com login, importacao CSV, assinaturas completas, relatorios PDF, API externa e responsividade mobile.",
        "Commits organizados por fase para demonstrar processo de desenvolvimento.",
    ]:
        story.append(bullet(item))

    story.append(p("10. Backlog recomendado", "H1Custom"))
    story.append(
        pdf_table(
            [
                ["Prioridade", "Item", "Motivo"],
                ["MVP", "Perfil, transacoes, dashboard, RealCost, metas e simulador", "Prova o diferencial com menor escopo viavel."],
                ["Pos-MVP", "Score financeiro e assinaturas completas", "Aumenta profundidade apos dados centrais estarem estaveis."],
                ["Pos-MVP", "Relatorios por periodo e exportacao PDF", "Ajuda apresentacao e evolucao tecnica."],
                ["Futuro", "Login e multiplos usuarios", "Necessario para produto real, mas nao para demonstracao inicial."],
                ["Futuro", "Importacao CSV e integracao bancaria", "Eleva complexidade e deve vir apenas depois da base validada."],
            ],
            [0.9 * inch, 2.55 * inch, 2.8 * inch],
        )
    )

    story.append(p("11. Regra de comunicacao do produto", "H1Custom"))
    story.append(
        p(
            "O FinanceScope deve ajudar o usuario a decidir, nao mandar nele. As mensagens devem ser objetivas, explicaveis e transparentes. Exemplo: 'Esta compra equivale a 18h40 de trabalho e pode atrasar sua meta em 1,2 mes, considerando os dados informados no perfil.'"
        )
    )
    for item in [
        "Use 'estimativa', 'pode indicar', 'com base nos seus dados' e 'no ritmo atual'.",
        "Evite frases como 'nao compre', 'voce deve investir' ou 'garantia de economia'.",
        "Mostre sempre a formula ou uma explicacao simples perto do resultado principal.",
    ]:
        story.append(bullet(item))

    story.append(p("12. Definicao de pronto do MVP", "H1Custom"))
    for idx, item in enumerate(
        [
            "Usuario configura perfil financeiro e visualiza dados no dashboard.",
            "Usuario cadastra, edita, filtra e remove transacoes.",
            "Dashboard mostra receitas, despesas, saldo, categorias, evolucao e alertas.",
            "RealCost Engine tem testes automatizados para os calculos principais.",
            "Usuario cria meta e o sistema calcula progresso e prazo estimado.",
            "Simulador mostra impacto de compra em horas, renda, risco e meta.",
            "Projeto possui README, dados demo, prints, aviso educacional e instrucoes de execucao.",
        ],
        start=1,
    ):
        story.append(p(f"{idx}. {item}"))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(0.75 * inch, 0.45 * inch, "FinanceScope - Plano de Desenvolvimento")
        canvas.drawRightString(7.75 * inch, 0.45 * inch, f"Pagina {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(f"DOCX: {DOCX_PATH}")
    print(f"PDF: {PDF_PATH}")


if __name__ == "__main__":
    main()
